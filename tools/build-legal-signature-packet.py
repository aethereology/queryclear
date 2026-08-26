"""Build the editable QueryClear legal/tax approval signature packet.

The default professional and management artifacts are clean execution copies;
the Board consent remains a controlled draft until the final director roster and
executed exhibits are available. Pass ``--draft`` to regenerate review copies.
"""

import argparse

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "legal" / "signature-packet"
OUT_DIR.mkdir(parents=True, exist_ok=True)

INK = RGBColor(31, 51, 64)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(92, 102, 112)
LIGHT = "F2F4F7"
LINE = "C9D0D7"
FONT = "Calibri"
TODAY = "August 25, 2026"

ENTITY = "SparkCreatives Inc."
BRAND = "QueryClear"
DOC_NO = "N25000003997"
EIN = "33-4477854"
ADDRESS = "6120 Caladesi Ct, Jacksonville, Florida 32258"


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run(run, *, size=11, bold=False, italic=False, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_run(run, size=9, color=MUTED)


def setup_document(title, subtitle, status="EXECUTION COPY"):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, before, after, color in (
        ("Title", 23, 0, 6, INK),
        ("Subtitle", 13, 0, 16, MUTED),
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, INK),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = FONT
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.10

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run(f"{BRAND} | Governance and Professional Approval Packet")
    set_run(hr, size=9, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_label = "Confidential execution copy" if status == "EXECUTION COPY" else "Confidential review draft"
    fr = fp.add_run(f"{footer_label} | Page ")
    set_run(fr, size=9, color=MUTED)
    add_page_field(fp)

    p = doc.add_paragraph(style="Title")
    set_run(p.add_run(title), size=23, bold=True, color=INK)
    p = doc.add_paragraph(style="Subtitle")
    set_run(p.add_run(subtitle), size=13, color=MUTED)

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2160, 7200])
    rows = [
        ("Entity", f"{ENTITY}, Florida not-for-profit corporation ({DOC_NO})"),
        ("Brand/activity", f"{BRAND} — commercial website, SEO, and AI-search-readiness services"),
        ("Prepared", TODAY),
        ("Status", status),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        shade_cell(row.cells[0], LIGHT)
        p1 = row.cells[0].paragraphs[0]
        set_run(p1.add_run(label), bold=True)
        p2 = row.cells[1].paragraphs[0]
        set_run(p2.add_run(value))
    doc.add_paragraph()
    return doc


def add_paragraph(doc, text, *, bold_prefix=None, italic=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), bold=True)
        set_run(p.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_run(p.add_run(text), italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run(p.add_run(item))


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run(p.add_run(item))


def add_field_lines(doc, fields):
    table = doc.add_table(rows=len(fields), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2600, 6760])
    for row, (label, value) in zip(table.rows, fields):
        shade_cell(row.cells[0], LIGHT)
        set_run(row.cells[0].paragraphs[0].add_run(label), bold=True)
        set_run(row.cells[1].paragraphs[0].add_run(value))
    return table


def add_signature(doc, signer_label, fields):
    doc.add_heading(signer_label, level=2)
    table = doc.add_table(rows=len(fields), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2600, 6760])
    for row, (label, value) in zip(table.rows, fields):
        shade_cell(row.cells[0], LIGHT)
        set_run(row.cells[0].paragraphs[0].add_run(label), bold=True)
        set_run(row.cells[1].paragraphs[0].add_run(value))
    add_paragraph(
        doc,
        "Electronic signatures and counterparts are permitted to the extent approved by the signer and the Corporation.",
        italic=True,
    )


def add_sources(doc, sources):
    doc.add_heading("Reference record", level=1)
    add_paragraph(
        doc,
        "These materials support review but do not substitute for the signer's independent professional analysis.",
        italic=True,
    )
    add_bullets(doc, sources)


def build_counsel(status="EXECUTION COPY"):
    doc = setup_document(
        "Counsel Legal-Structure Review and Concurrence",
        "QueryClear commercial activity operated under SparkCreatives Inc.",
        status,
    )
    add_paragraph(
        doc,
        "Purpose. This memorandum creates a clear written record of counsel's independent review of the Corporation's authority and the legal structure selected for QueryClear. It is not management's legal opinion. Counsel should revise any statement that does not match counsel's analysis before signing.",
        bold_prefix="Purpose.",
    )
    doc.add_heading("1. Facts presented for review", level=1)
    add_bullets(doc, [
        f"{ENTITY} is an active Florida not-for-profit corporation, document number {DOC_NO}, with principal address {ADDRESS}.",
        f"The Internal Revenue Service has recognized the Corporation as an organization described in section 501(c)(3); EIN {EIN}.",
        f"{BRAND} is presented publicly as a brand or trade name of the Corporation and offers paid SEO, website, audit, and recurring care services.",
        "Public pricing presently includes a no-cost automated preview, a $497 human-delivered audit, implementation beginning at $2,500, website builds beginning at $6,500, and an eligibility-gated $997 monthly care service.",
        "Service fees are intended to be payments for services and are not intended to be solicited, receipted, or represented as charitable contributions.",
        "Management has delivered the separately signed Management Factual Representations and Implementation Certification and disclosed material shared personnel, systems, intellectual property, compensation, customer-acquisition practices, and data flows.",
    ])
    doc.add_heading("2. Scope of counsel review", level=1)
    add_numbered(doc, [
        "Review the articles, bylaws, IRS determination materials, exemption application and supporting narrative, board records, conflicts policy, and relevant agreements.",
        "Determine whether and on what conditions the QueryClear activity may be operated directly by the Corporation consistently with its governing documents and exempt purposes.",
        "Evaluate private inurement, private benefit, excess-benefit, compensation, conflict-of-interest, donated-asset, and related-party concerns.",
        "Advise whether a taxable subsidiary, separate taxable operator, cost-sharing arrangement, license, or other structural separation is required or prudent.",
        "Confirm required Florida fictitious-name, local licensing, contract, consumer, privacy, advertising, email, and telemarketing compliance steps within the agreed scope.",
        "Review the board written consent, online service terms, MSA/SOW, DPA, privacy disclosures, and outbound-marketing controls for material legal sufficiency.",
    ])
    doc.add_heading("3. Counsel's selected conclusion", level=1)
    add_paragraph(doc, "Counsel must select one conclusion and strike the others before execution.", italic=True)
    add_field_lines(doc, [
        ("Conclusion A", "[  ] Direct operation by the Corporation is approved, subject to the conditions in Section 4."),
        ("Conclusion B", "[  ] Operation through a taxable subsidiary or affiliate is required or recommended; structure and effective date are stated in Section 4."),
        ("Conclusion C", "[  ] Another structure is approved; the structure and rationale are stated in Section 4."),
        ("Launch status", "[  ] Approved now   [  ] Approved after listed conditions   [  ] Not approved"),
    ])
    doc.add_heading("4. Required conditions and rationale", level=1)
    add_field_lines(doc, [
        ("Structure", "________________________________________________________________________________"),
        ("Exempt-purpose analysis", "________________________________________________________________________________"),
        ("Private-benefit / compensation controls", "________________________________________________________________________________"),
        ("Board / conflict actions", "________________________________________________________________________________"),
        ("Trade name / licensing", "________________________________________________________________________________"),
        ("Contracts / privacy / marketing", "________________________________________________________________________________"),
        ("Other launch conditions", "________________________________________________________________________________"),
        ("Required review date", "________________________________________________________________________________"),
    ])
    doc.add_heading("5. Counsel concurrence", level=1)
    add_paragraph(
        doc,
        "Based on the documents and facts identified above, and subject to the assumptions, qualifications, jurisdictional limits, and conditions written into this memorandum, the undersigned counsel concurs with the selected conclusion in Section 3. Counsel has modified this document wherever necessary so that it accurately states counsel's professional conclusion. No unsigned draft is an approval or legal opinion.",
    )
    add_signature(doc, "Attorney signature", [
        ("Attorney signature", "________________________________________________   Date: __________________"),
        ("Printed name", "________________________________________________________________________________"),
        ("Firm", "________________________________________________________________________________"),
        ("Bar / jurisdiction", "________________________________________________________________________________"),
        ("Matter / file number", "________________________________________________________________________________"),
    ])
    add_sources(doc, [
        "Florida Division of Corporations record for SparkCreatives Inc., document N25000003997.",
        "IRS determination letter and exemption materials maintained by the Corporation.",
        "IRS, Unrelated Business Income Defined; Unrelated Business Income Tax; Substantially Related; Inurement/Private Benefit guidance.",
        "Current QueryClear website, offer descriptions, Privacy Policy, Terms, checkout, outreach, and service-delivery materials.",
    ])
    path = OUT_DIR / "01-counsel-legal-structure-concurrence.docx"
    doc.save(path)
    return path


def build_cpa(status="EXECUTION COPY"):
    doc = setup_document(
        "CPA Tax and Accounting Position Concurrence",
        "QueryClear revenue, expense allocation, reporting, and ongoing controls",
        status,
    )
    add_paragraph(
        doc,
        "Purpose. This memorandum documents the tax and accounting treatment selected for QueryClear after the CPA's independent review. It must be conformed to the legal structure approved by counsel. The CPA should revise all positions, thresholds, filing forms, and effective dates before signing.",
        bold_prefix="Purpose.",
    )
    doc.add_heading("1. Facts and records supplied", level=1)
    add_bullets(doc, [
        f"Taxpayer: {ENTITY}; EIN {EIN}; Florida document number {DOC_NO}.",
        f"Activity: paid services sold under the {BRAND} name, including audits, website work, recurring care, and related digital services.",
        "Revenue may be collected through Stripe or invoicing; operating costs may include software, AI/API usage, hosting, payment fees, contractors, payroll, marketing, insurance, and shared overhead.",
        "Management will maintain separate activity-level accounting and will not characterize commercial service payments as charitable gifts.",
        "Counsel's executed legal-structure concurrence and management's signed factual representations are incorporated by reference.",
    ])
    doc.add_heading("2. Positions to be completed by the CPA", level=1)
    add_field_lines(doc, [
        ("Legal operating model", "________________________________________________________________________________"),
        ("Revenue classification", "[  ] Related program service   [  ] Unrelated business   [  ] Taxable affiliate   [  ] Other"),
        ("Federal filings", "[  ] Form 990   [  ] Form 990-T   [  ] Form 1120   [  ] Estimated tax   [  ] Other"),
        ("State / local filings", "________________________________________________________________________________"),
        ("Sales-tax position", "________________________________________________________________________________"),
        ("Nexus review scope", "________________________________________________________________________________"),
        ("Accounting basis", "[  ] Cash   [  ] Accrual   [  ] Other: ______________________________________________"),
        ("Effective date", "________________________________________________________________________________"),
    ])
    doc.add_heading("3. Required accounting controls", level=1)
    add_numbered(doc, [
        "Create a distinct QueryClear class, department, fund, or entity ledger and reconcile Stripe, bank, invoicing, refunds, disputes, and accounts receivable monthly.",
        "Use a written, consistently applied method to allocate shared labor, software, hosting, facilities, marketing, contractor, and administrative expenses.",
        "Document compensation and related-party payments against board-approved arrangements and contemporaneous support.",
        "Separate charitable contributions, grants, restricted funds, in-kind donations, and volunteer activity from commercial service receipts and costs.",
        "Retain contracts, invoices, source records, customer refunds, processor reports, allocation workpapers, and tax support under the approved retention schedule.",
        "Review revenue mix, gross unrelated-business income, net taxable income, estimated-tax needs, nexus, and filing thresholds at least quarterly during the first operating year.",
    ])
    doc.add_heading("4. CPA-specific instructions and thresholds", level=1)
    add_field_lines(doc, [
        ("Chart-of-accounts instructions", "________________________________________________________________________________"),
        ("Allocation methodology", "________________________________________________________________________________"),
        ("Quarterly review triggers", "________________________________________________________________________________"),
        ("Estimated-tax instructions", "________________________________________________________________________________"),
        ("Form 990 presentation", "________________________________________________________________________________"),
        ("Intercompany treatment", "________________________________________________________________________________"),
        ("Record-retention additions", "________________________________________________________________________________"),
        ("Next formal review", "________________________________________________________________________________"),
    ])
    doc.add_heading("5. CPA concurrence", level=1)
    add_paragraph(
        doc,
        "Based on the records and representations identified above, and subject to the assumptions and qualifications written into this memorandum, the undersigned CPA concurs with the completed tax and accounting positions. The CPA has modified this document wherever necessary so it accurately states the selected treatment. No unsigned draft is tax advice, a filing position, or an attestation engagement.",
    )
    add_signature(doc, "CPA signature", [
        ("CPA signature", "________________________________________________   Date: __________________"),
        ("Printed name", "________________________________________________________________________________"),
        ("Firm", "________________________________________________________________________________"),
        ("License / jurisdiction", "________________________________________________________________________________"),
        ("Engagement / file number", "________________________________________________________________________________"),
    ])
    add_sources(doc, [
        "Executed counsel legal-structure concurrence and the draft board written consent.",
        "IRS determination letter, exemption application, governing documents, and prior Form 990-series filings, if any.",
        "QueryClear revenue, Stripe, bank, payroll, contractor, software, marketing, shared-cost, and customer-refund records.",
        "IRS guidance and current filing instructions applicable to the CPA's selected treatment and tax year.",
    ])
    path = OUT_DIR / "02-cpa-tax-accounting-concurrence.docx"
    doc.save(path)
    return path


def build_board():
    doc = setup_document(
        "Unanimous Written Consent of the Board of Directors",
        "Ratification and governance of the QueryClear commercial activity",
        status="DRAFT — COMPLETE AFTER COUNSEL AND CPA CONCURRENCES",
    )
    add_paragraph(
        doc,
        "The undersigned, constituting all directors of SparkCreatives Inc. entitled to vote on these actions, adopt the following resolutions by unanimous written consent effective on the date stated below. The Secretary shall place the executed consent and incorporated exhibits in the Corporation's minute book.",
    )
    doc.add_heading("Recitals", level=1)
    add_bullets(doc, [
        f"The Corporation operates or proposes to operate paid services using the trade name {BRAND}.",
        "The Board has reviewed management's factual representations, counsel's legal-structure concurrence, the CPA's tax/accounting concurrence, projected budgets, compensation and conflict disclosures, and the implementation controls listed below.",
        "The Board intends to preserve the Corporation's exempt purposes, prevent private inurement or impermissible private benefit, maintain accurate books and records, and ensure that service fees are not represented as charitable contributions.",
    ])
    doc.add_heading("Resolutions", level=1)
    add_numbered(doc, [
        "The Board approves QueryClear only under the legal operating structure, conditions, limitations, and effective date stated in the executed counsel concurrence.",
        "The Board adopts the revenue classification, entity/fund accounting, expense-allocation method, filing treatment, review triggers, and effective date stated in the executed CPA concurrence.",
        "The officers are authorized to register and maintain the QueryClear fictitious name and applicable licenses; align website, contracts, W-9, invoices, bank and Stripe records; and execute approved commercial agreements.",
        "Commercial service receipts shall be invoiced and recorded as service revenue and shall not be solicited, acknowledged, or receipted as tax-deductible charitable contributions.",
        "All shared personnel, intellectual property, software, equipment, facilities, marketing, and administrative expenses shall be tracked and allocated under the CPA-approved policy, with any intercompany or related-party arrangements documented in writing and on arm's-length terms.",
        "Any officer, director, employee, or related person with a financial interest shall disclose it and abstain as required. Compensation and related-party transactions require review by disinterested directors using appropriate comparability information and contemporaneous minutes.",
        "Management shall use approved online service terms for self-service products, an MSA and SOW for implementation work, a DPA/security addendum where applicable, accurate privacy and marketing disclosures, and documented refund/cancellation procedures.",
        "Management shall maintain durable lead and payment records, secure client reports, least-privilege access, incident response, vendor review, data retention, and documented service-delivery controls.",
        "Before scaled promotion, management shall confirm professional insurance, proof substantiation, audit-methodology accuracy, fulfillment capacity, and completion of every launch condition identified by counsel and the CPA.",
        "The Board ratifies prior good-faith organizational acts concerning QueryClear only to the extent consistent with these resolutions and directs management to remediate any inconsistency promptly.",
        "The Board will review QueryClear at least annually and sooner upon a material change in pricing, revenue, services, ownership, compensation, jurisdictions, data processing, or tax/legal guidance.",
    ])
    doc.add_heading("Incorporated exhibits", level=1)
    add_bullets(doc, [
        "Exhibit A — Executed Counsel Legal-Structure Review and Concurrence.",
        "Exhibit B — Executed CPA Tax and Accounting Position Concurrence.",
        "Exhibit C — Executed Management Factual Representations and Implementation Certification.",
        "Exhibit D — Approved shared-resource/cost-allocation schedule and compensation/conflict record.",
    ])
    doc.add_heading("Effective date and signatures", level=1)
    add_field_lines(doc, [
        ("Effective date", "________________________________________________________________________________"),
        ("Meeting / consent record", "________________________________________________________________________________"),
        ("Total directors entitled to vote", "________________________________________________________________________________"),
    ])
    add_paragraph(
        doc,
        "Every director entitled to vote must sign. Duplicate this signature page if needed and strike unused blocks before the Secretary certifies the final packet.",
        italic=True,
    )
    for index in range(1, 7):
        if index in (3, 5):
            doc.add_page_break()
        add_signature(doc, f"Director signature block {index}", [
            ("Signature", "________________________________________________   Date: __________________"),
            ("Printed name", "________________________________________________________________________________"),
            ("Conflict disclosed?", "[  ] No   [  ] Yes — disclosure/abstention record attached"),
        ])
    doc.add_page_break()
    add_signature(doc, "Secretary certification", [
        ("Secretary signature", "________________________________________________   Date: __________________"),
        ("Printed name", "________________________________________________________________________________"),
        ("Certification", "I certify that the above signatures constitute all directors entitled to vote and that this consent is retained in the corporate records."),
    ])
    path = OUT_DIR / "03-board-unanimous-written-consent.docx"
    doc.save(path)
    return path


def build_management(status="EXECUTION COPY"):
    doc = setup_document(
        "Management Factual Representations and Implementation Certification",
        "Facts supplied to counsel, CPA, and the Board concerning QueryClear",
        status,
    )
    add_paragraph(
        doc,
        "Purpose. Counsel and the CPA may rely on this document only to the extent stated in their respective engagements. Management must correct, supplement, or qualify every item that is incomplete or inaccurate before signing.",
        bold_prefix="Purpose.",
    )
    doc.add_heading("1. Corporate and commercial facts", level=1)
    add_bullets(doc, [
        f"The legal entity presented as operating QueryClear is {ENTITY}, Florida document number {DOC_NO}, EIN {EIN}, with principal address {ADDRESS}.",
        "QueryClear is not presently represented to customers as a separate incorporated entity. Management will identify any other seller, subsidiary, affiliate, DBA registration, or merchant account in the exceptions schedule.",
        "QueryClear provides or proposes to provide paid audits, website optimization, website builds, recurring care, and related services; current public pricing is disclosed in the counsel memorandum.",
        "Service fees are intended to be commercial payments, not charitable gifts, and management will not issue charitable contribution acknowledgments for them.",
        "Management has disclosed all current and proposed compensation, contractors, related parties, donated assets, shared personnel, shared systems, and intellectual property used by QueryClear.",
    ])
    doc.add_heading("2. Customer, data, and marketing facts", level=1)
    add_bullets(doc, [
        "The system may process lead identity and contact data, website URLs and crawled public content, technical logs and IP addresses, payment and subscription metadata, attribution fields, prospect data from public sources, email delivery/engagement records, report-view events, and client-access credentials when separately authorized.",
        "Current or contemplated providers include Vercel, Cloudflare, Upstash, Resend, Stripe, Microsoft, Apify, OpenAI or another audit runtime/model provider, and other processors disclosed in the maintained vendor register.",
        "Management will not knowingly solicit patient records, medical records, passwords, payment-card details, or other unnecessary sensitive data through public forms.",
        "All public claims, AI-test labels, case studies, testimonials, logos, credentials, reviews, and outcome metrics will be supported by retained evidence and required permissions. Fictional demonstrations will remain clearly labeled.",
        "Cold outreach, telephone activity, tracking, and follow-up cadences will be operated only under the approved jurisdictional, disclosure, suppression, and recordkeeping procedures.",
    ])
    doc.add_heading("3. Exceptions and corrections", level=1)
    add_field_lines(doc, [
        ("Other seller / affiliate / DBA", "None, except: _____________________________________________________________________"),
        ("Related parties / compensation", "None, except: _____________________________________________________________________"),
        ("Shared or donated assets", "None, except: _____________________________________________________________________"),
        ("Outstanding tax / filing issues", "None known, except: ______________________________________________________________"),
        ("Claims or disputes", "None known, except: ______________________________________________________________"),
        ("Data / security incidents", "None known, except: ______________________________________________________________"),
        ("Other material qualification", "________________________________________________________________________________"),
    ])
    doc.add_heading("4. Implementation certification", level=1)
    add_paragraph(doc, "Management shall initial each completed control and leave incomplete controls uninitialed.", italic=True)
    controls = [
        "Counsel concurrence executed and all launch conditions completed.",
        "CPA concurrence executed; ledger, allocations, filings, and review calendar configured.",
        "Board written consent executed and retained in the minute book.",
        "QueryClear fictitious name and required business licenses verified or registered.",
        "Legal seller aligned across website, Terms, Privacy, contracts, W-9, invoices, Stripe, bank, and email.",
        "Online service terms use recorded clickwrap; MSA/SOW and DPA/security addendum are in use as applicable.",
        "Refund, cancellation, chargeback, subscription, and customer-portal paths were tested end to end.",
        "Leads and Stripe events are stored durably with idempotency, retries, reconciliation, and alerts.",
        "Real client reports require authenticated or signed expiring access and are not publicly cached.",
        "Public audit claims match the documented live/modelled methodology; URL fetching is SSRF-hardened.",
        "Privacy notice, retention schedule, subprocessor list, incident response, access control, and outbound procedures are current.",
        "Professional/E&O and cyber insurance requirements were completed or expressly waived by the Board after advice.",
        "Case-study/testimonial evidence and releases are retained; no fictional proof is presented as real.",
        "Service scopes, capacity, QA, staging, backup, rollback, analytics, handoff, and support SOPs are operational.",
    ]
    table = doc.add_table(rows=len(controls) + 1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [900, 7080, 1380])
    headers = ["Initials", "Control", "Date"]
    for idx, header in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], LIGHT)
        set_run(table.rows[0].cells[idx].paragraphs[0].add_run(header), bold=True)
    for row, control in zip(table.rows[1:], controls):
        set_run(row.cells[0].paragraphs[0].add_run("________"))
        set_run(row.cells[1].paragraphs[0].add_run(control))
        set_run(row.cells[2].paragraphs[0].add_run("__________"))
    doc.add_heading("5. Management certification", level=1)
    add_paragraph(
        doc,
        "The undersigned certify that they reviewed this document, made reasonable inquiry, disclosed all material facts known to them, and will promptly notify counsel, the CPA, and the Board of any material correction or change. This certification does not replace the independent duties of any professional or director.",
    )
    add_signature(doc, "Authorized officer", [
        ("Signature", "________________________________________________   Date: __________________"),
        ("Printed name / title", "________________________________________________________________________________"),
    ])
    add_signature(doc, "QueryClear operating lead", [
        ("Signature", "________________________________________________   Date: __________________"),
        ("Printed name / title", "________________________________________________________________________________"),
    ])
    path = OUT_DIR / "04-management-representations-and-implementation-certification.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--draft",
        action="store_true",
        help="Generate review drafts instead of professional execution copies.",
    )
    args = parser.parse_args()
    status = "DRAFT FOR PROFESSIONAL REVIEW" if args.draft else "EXECUTION COPY"
    created = [
        build_counsel(status),
        build_cpa(status),
        build_board(),
        build_management(status),
    ]
    for item in created:
        print(item)
